"""
Business OS generator for the "Planning Control Tower" venture.

Produces:
  1. Weighted opportunity scoring table (markdown)      -> scoring_table.md
  2. 3-scenario 12-month financial model (markdown)     -> financial_model.md
  3. Business_OS.xlsx with live formulas:
       - Assumptions, Base/Conservative/Aggressive model
       - Opportunity scoring, CRM pipeline, 90-day tracker, KPI dashboard

All money in PKR unless stated. FX and every driver sit on the Assumptions
sheet so the model can be re-run with your own numbers.

Usage:  python build_model.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent

# ---------------------------------------------------------------- assumptions
FX = 280               # PKR per USD (ASSUMPTION - update monthly)
UPWORK_FEE = 0.10      # Upwork service fee on Upwork-sourced revenue (ASSUMPTION: variable 0-15%)
PAYMENT_FEE = 0.02     # Payoneer / transfer / FX spread on all revenue (ASSUMPTION)
SETUP_COST_PKR = 60_000  # one-off in Month 1 (domain, site, NTN/PSEB, demo data, design)

# ------------------------------------------------------------ scoring weights
WEIGHTS = {
    "Market demand": 0.10, "Profitability": 0.10, "Low capital": 0.07,
    "Ease of start": 0.08, "Skill fit": 0.15, "Intl earning": 0.12,
    "Scalability": 0.10, "Competition (10=low)": 0.08,
    "Speed to revenue": 0.10, "Long-term potential": 0.10,
}
assert abs(sum(WEIGHTS.values()) - 1) < 1e-9

IDEAS = [
    # name, scores in WEIGHTS order
    ("1. Planning & Plant Performance Control Tower (productized Power BI for SME manufacturers)", [8, 8, 9, 7, 10, 9, 7, 6, 7, 8]),
    ("2. Fractional S&OP / Supply Planning Analyst (retainer)",            [7, 8, 10, 6, 10, 8, 5, 7, 6, 7]),
    ("3. Inventory & SKU Health Analytics for DTC/Amazon brands",          [8, 7, 9, 6, 8, 9, 7, 5, 6, 7]),
    ("4. Excel / Power Query reporting automation (productized gigs)",     [8, 6, 10, 9, 9, 8, 5, 3, 9, 5]),
    ("5. Power BI / Excel KPI template store (digital products)",          [6, 5, 9, 7, 9, 9, 9, 4, 5, 6]),
    ("6. OEE & downtime micro-SaaS for small factories",                   [6, 8, 6, 4, 8, 7, 9, 5, 3, 8]),
    ("7. Rough-cut capacity planner micro-SaaS for job shops",             [5, 8, 6, 3, 9, 7, 9, 6, 2, 7]),
    ("8. KPI / Balanced Scorecard office for GCC SMEs",                    [7, 8, 10, 5, 10, 8, 4, 6, 5, 7]),
    ("9. AI document extraction for freight forwarders",                   [8, 8, 9, 6, 6, 9, 7, 4, 6, 6]),
    ("10. ESG & carbon data reporting for South-Asian textile exporters",  [7, 7, 9, 5, 7, 5, 6, 7, 4, 8]),
    ("11. Remote SAP PP / IBP functional support",                         [7, 7, 10, 6, 8, 8, 3, 5, 6, 6]),
    ("12. White-label BI delivery for ERP partners & BI agencies",         [8, 6, 10, 6, 9, 9, 6, 6, 6, 6]),
    ("13. AI-assisted market & competitor research reports",               [5, 5, 10, 7, 8, 8, 4, 3, 7, 3]),
    ("14. Cohort course: Power BI for Supply Chain Planners",              [6, 8, 9, 5, 9, 8, 8, 5, 4, 7]),
    ("15. AI order-entry automation for distributors (email PO -> ERP)",   [8, 8, 9, 5, 7, 9, 7, 5, 5, 7]),
    ("16. Amazon/Shopify SKU profitability dashboards",                    [7, 6, 9, 6, 7, 9, 6, 3, 6, 5]),
    ("17. S/4HANA migration master-data readiness (subcontract to SIs)",   [8, 8, 10, 4, 9, 8, 5, 6, 4, 7]),
    ("18. Pakistan sourcing intelligence for foreign buyers",              [5, 6, 8, 4, 5, 8, 4, 7, 4, 5]),
    ("19. Remote supply-chain analytics talent marketplace",               [5, 7, 7, 3, 6, 8, 8, 6, 2, 7]),
    ("20. Spreadsheet MRP/BOM toolkit for small makers",                   [6, 5, 9, 7, 8, 9, 8, 4, 6, 5]),
]

# -------------------------------------------------------------- scenarios
# Per month (1..12): projects, avg project USD, retainers, avg retainer USD,
# share of revenue via Upwork, contractor PKR, software USD, marketing USD
SCENARIOS = {
    "Conservative": dict(
        projects=[0, 0, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1],
        proj_usd=[0, 0, 350, 450, 500, 550, 600, 600, 650, 650, 700, 700],
        retainers=[0, 0, 0, 0, 1, 1, 1, 1, 2, 2, 2, 2],
        ret_usd=[0, 0, 0, 0, 400, 400, 400, 400, 425, 425, 450, 450],
        upwork=[1, 1, 1, 1, .9, .9, .8, .8, .7, .7, .7, .7],
        contractor=[0] * 12,
        software=[80, 100, 100, 100, 100, 100, 100, 100, 100, 100, 100, 100],
        marketing=[40, 60, 60, 60, 60, 60, 60, 60, 60, 60, 60, 60],
    ),
    "Base": dict(
        projects=[0, 1, 1, 2, 2, 2, 2, 2, 2, 2, 2, 2],
        proj_usd=[0, 400, 600, 650, 750, 800, 850, 900, 950, 1000, 1050, 1100],
        retainers=[0, 0, 0, 0, 1, 1, 2, 2, 3, 3, 4, 4],
        ret_usd=[0, 0, 0, 0, 500, 500, 550, 600, 600, 650, 650, 650],
        upwork=[1, 1, 1, .9, .8, .7, .6, .6, .5, .5, .4, .4],
        contractor=[0, 0, 0, 0, 0, 0, 60_000, 60_000, 80_000, 80_000, 100_000, 100_000],
        software=[90, 140, 140, 140, 150, 150, 150, 150, 160, 160, 160, 160],
        marketing=[60, 160, 160, 160, 180, 180, 180, 180, 200, 200, 200, 200],
    ),
    "Aggressive": dict(
        projects=[0, 1, 2, 3, 3, 3, 3, 3, 3, 3, 3, 3],
        proj_usd=[0, 600, 800, 900, 1000, 1100, 1200, 1300, 1400, 1400, 1500, 1500],
        retainers=[0, 0, 1, 1, 2, 3, 4, 5, 5, 6, 7, 8],
        ret_usd=[0, 0, 500, 550, 600, 650, 700, 700, 750, 750, 800, 800],
        upwork=[1, .9, .8, .7, .6, .5, .5, .4, .4, .3, .3, .3],
        contractor=[0, 0, 0, 0, 60_000, 80_000, 120_000, 150_000, 180_000, 220_000, 260_000, 300_000],
        software=[100, 180, 180, 200, 200, 220, 220, 250, 250, 250, 280, 280],
        marketing=[80, 200, 250, 250, 300, 300, 350, 350, 400, 400, 400, 400],
    ),
}


def weighted(scores):
    return round(sum(s * w for s, w in zip(scores, WEIGHTS.values())), 2)


def scoring_md():
    hdr = ["Idea", *[k.split(" (")[0] for k in WEIGHTS], "Weighted"]
    rows = sorted(IDEAS, key=lambda x: -weighted(x[1]))
    out = ["| Rank | " + " | ".join(hdr) + " |", "|" + "---|" * (len(hdr) + 1)]
    for i, (n, s) in enumerate(rows, 1):
        out.append(f"| {i} | {n} | " + " | ".join(map(str, s)) + f" | **{weighted(s):.2f}** |")
    w = " · ".join(f"{k} {int(v*100)}%" for k, v in WEIGHTS.items())
    return f"Weights: {w}\n\n" + "\n".join(out)


def run(sc):
    rows, cum_cost, cum_profit, cum_inv = [], 0, 0, 0
    for m in range(12):
        rev_usd = sc["projects"][m] * sc["proj_usd"][m] + sc["retainers"][m] * sc["ret_usd"][m]
        cust = sc["projects"][m] + sc["retainers"][m]
        rev = rev_usd * FX
        direct = rev * (UPWORK_FEE * sc["upwork"][m] + PAYMENT_FEE) + sc["contractor"][m]
        soft = sc["software"][m] * FX + (SETUP_COST_PKR if m == 0 else 0)
        mkt = sc["marketing"][m] * FX
        net = rev - direct - soft - mkt
        cum_profit += net
        # cumulative investment = deepest cumulative cash hole funded from capital
        cum_inv = max(cum_inv, -cum_profit)
        rows.append(dict(m=m + 1, cust=cust, arpc=(rev / cust if cust else 0), rev=rev,
                         direct=direct, soft=soft, mkt=mkt, net=net,
                         inv=cum_inv, cum=cum_profit))
    return rows


def k(x):
    return f"{x/1000:,.0f}k"


def financial_md():
    parts = [f"FX assumption: PKR {FX}/USD · Upwork fee {int(UPWORK_FEE*100)}% on Upwork-sourced revenue · "
             f"payment/FX cost {int(PAYMENT_FEE*100)}% · one-off setup PKR {SETUP_COST_PKR:,} in Month 1 (shown in Software/Ops).\n"
             "Cumulative investment = peak capital actually drawn from your PKR 500k (deepest cumulative cash hole).\n"]
    for name, sc in SCENARIOS.items():
        r = run(sc)
        parts.append(f"\n#### {name} scenario (PKR)\n")
        parts.append("| Month | Paying customers | Avg rev / customer | Revenue | Direct costs | Software / ops | Marketing | Net profit | Cumulative investment | Cumulative profit |")
        parts.append("|---|---|---|---|---|---|---|---|---|---|")
        for x in r:
            if x["m"] in (1, 2, 3, 4, 5, 6, 12):
                parts.append(f"| M{x['m']} | {x['cust']} | {k(x['arpc'])} | {k(x['rev'])} | {k(x['direct'])} | {k(x['soft'])} | "
                             f"{k(x['mkt'])} | {k(x['net'])} | {k(x['inv'])} | {k(x['cum'])} |")
        tot = sum(x["rev"] for x in r)
        parts.append(f"\n12-month revenue: **PKR {tot/1e6:,.2f}M** (≈ USD {tot/FX:,.0f}) · "
                     f"Month-12 run-rate: **PKR {r[-1]['rev']/1e6:,.2f}M/month** · "
                     f"peak capital used: **PKR {k(max(x['inv'] for x in r))}** · "
                     f"12-month cumulative net: **PKR {k(r[-1]['cum'])}**")
    return "\n".join(parts)


# ----------------------------------------------------------------- workbook
H = Font(bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F3864")
INP = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def header(ws, row, values, widths=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font, cell.fill = H, HF
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def build_xlsx():
    wb = Workbook()

    # Assumptions
    a = wb.active
    a.title = "Assumptions"
    header(a, 1, ["Driver", "Value", "Note"], [34, 14, 70])
    rows = [("FX PKR per USD", FX, "ASSUMPTION - update monthly from SBP interbank"),
            ("Upwork fee %", UPWORK_FEE, "ASSUMPTION - Upwork variable fee; check your contract"),
            ("Payment / FX cost %", PAYMENT_FEE, "Payoneer withdrawal + spread, ASSUMPTION"),
            ("One-off setup PKR (Month 1)", SETUP_COST_PKR, "Domain, site, NTN/PSEB, demo data, design"),
            ("Starting capital PKR", 500_000, "Hard ceiling - do not exceed")]
    for i, (n, v, note) in enumerate(rows, 2):
        a.cell(row=i, column=1, value=n)
        c = a.cell(row=i, column=2, value=v)
        c.fill = INP
        a.cell(row=i, column=3, value=note)
    a["B3"].number_format = a["B4"].number_format = "0%"

    # Scenario sheets with formulas
    cols = ["Month", "Projects", "Avg project USD", "Retainers", "Avg retainer USD", "Upwork share",
            "Contractor PKR", "Software USD", "Marketing USD", "Paying customers", "Revenue PKR",
            "Avg rev/customer PKR", "Direct costs PKR", "Software/ops PKR", "Marketing PKR",
            "Net profit PKR", "Cumulative profit PKR", "Cumulative investment PKR"]
    for name, sc in SCENARIOS.items():
        ws = wb.create_sheet(name)
        header(ws, 1, cols, [8] + [13] * (len(cols) - 1))
        for m in range(12):
            r = m + 2
            vals = [m + 1, sc["projects"][m], sc["proj_usd"][m], sc["retainers"][m], sc["ret_usd"][m],
                    sc["upwork"][m], sc["contractor"][m], sc["software"][m], sc["marketing"][m]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                if c > 1:
                    cell.fill = INP
            f = {
                10: f"=B{r}+D{r}",
                11: f"=(B{r}*C{r}+D{r}*E{r})*Assumptions!$B$2",
                12: f"=IFERROR(K{r}/J{r},0)",
                13: f"=K{r}*(Assumptions!$B$3*F{r}+Assumptions!$B$4)+G{r}",
                14: f"=H{r}*Assumptions!$B$2+IF(A{r}=1,Assumptions!$B$5,0)",
                15: f"=I{r}*Assumptions!$B$2",
                16: f"=K{r}-M{r}-N{r}-O{r}",
                17: f"=P{r}" if m == 0 else f"=Q{r-1}+P{r}",
                18: f"=MAX(0,-Q{r})" if m == 0 else f"=MAX(R{r-1},-Q{r})",
            }
            for c, fx in f.items():
                ws.cell(row=r, column=c, value=fx).number_format = "#,##0"
            ws.cell(row=r, column=6).number_format = "0%"
        ws.cell(row=15, column=1, value="Total").font = Font(bold=True)
        for c in (11, 13, 14, 15, 16):
            L = get_column_letter(c)
            ws.cell(row=15, column=c, value=f"=SUM({L}2:{L}13)").number_format = "#,##0"
        ws.freeze_panes = "B2"

    # Opportunity scoring
    s = wb.create_sheet("Opportunity Scoring")
    header(s, 1, ["Idea", *WEIGHTS.keys(), "Weighted score"], [70] + [11] * 11)
    s.cell(row=2, column=1, value="WEIGHTS").font = Font(bold=True)
    for c, w in enumerate(WEIGHTS.values(), 2):
        s.cell(row=2, column=c, value=w).fill = INP
    for i, (n, sc) in enumerate(IDEAS, 3):
        s.cell(row=i, column=1, value=n)
        for c, v in enumerate(sc, 2):
            s.cell(row=i, column=c, value=v).fill = INP
        s.cell(row=i, column=12, value=f"=SUMPRODUCT(B{i}:K{i},$B$2:$K$2)").number_format = "0.00"

    # CRM pipeline
    crm = wb.create_sheet("CRM Pipeline")
    crm_cols = ["Lead ID", "Company", "Country", "Industry", "Revenue band", "ERP", "Contact name", "Title",
                "LinkedIn URL", "Email", "Source", "Trigger / hook", "Stage", "First touch date",
                "Last touch date", "Next action", "Next action date", "Touches", "Deal value USD",
                "Probability", "Weighted USD", "Notes"]
    header(crm, 1, crm_cols, [8, 22, 10, 16, 12, 12, 18, 18, 28, 26, 12, 30, 16, 12, 12, 26, 12, 8, 12, 10, 12, 30])
    stages = ["1-Identified", "2-Contacted", "3-Replied", "4-Discovery booked", "5-Diagnostic proposed",
              "6-Won", "7-Lost", "8-Nurture"]
    dv = DataValidation(type="list", formula1='"' + ",".join(stages) + '"', allow_blank=True)
    crm.add_data_validation(dv)
    dv.add("M2:M1000")
    src = DataValidation(type="list", formula1='"LinkedIn,Cold email,Upwork,Partner,Referral,Community,Inbound"', allow_blank=True)
    crm.add_data_validation(src)
    src.add("K2:K1000")
    prob = {"1-Identified": 0.02, "2-Contacted": 0.05, "3-Replied": 0.15, "4-Discovery booked": 0.3,
            "5-Diagnostic proposed": 0.5, "6-Won": 1, "7-Lost": 0, "8-Nurture": 0.03}
    pm = crm.parent.create_sheet("Lists")
    for i, (st, p) in enumerate(prob.items(), 1):
        pm.cell(row=i, column=1, value=st)
        pm.cell(row=i, column=2, value=p)
    for r in range(2, 301):
        crm.cell(row=r, column=20, value=f'=IFERROR(VLOOKUP(M{r},Lists!$A$1:$B$8,2,FALSE),"")').number_format = "0%"
        crm.cell(row=r, column=21, value=f'=IFERROR(S{r}*T{r},"")').number_format = "#,##0"
    crm.cell(row=2, column=1, value="L001")
    crm.cell(row=2, column=2, value="(example) Acme Fasteners Ltd")
    crm.cell(row=2, column=13, value="2-Contacted")
    crm.cell(row=2, column=19, value=1200)
    crm.freeze_panes = "C2"

    # 90-day tracker
    t = wb.create_sheet("90-Day Tracker")
    header(t, 1, ["Week", "Phase", "Outcome target", "Prospects added", "Touches sent", "Replies",
                  "Discovery calls", "Proposals", "Wins", "Revenue USD", "Hours spent", "Status", "Blocker / learning"],
           [6, 18, 44, 10, 10, 9, 10, 10, 7, 11, 10, 12, 40])
    phases = (["Foundation"] * 1 + ["Acquisition system"] * 3 + ["First customers"] * 5 + ["Scale & productize"] * 4)
    targets = ["Offer, demo, site, payments live; 50-lead list", "100 touches; 3 Upwork proposals/day",
               "150 touches; 2 discovery calls", "150 touches; first proposal", "First paid diagnostic",
               "Deliver + testimonial", "Case study #1 published", "2nd/3rd client; raise price 20%",
               "First retainer signed", "SOPs v1 + contractor trial", "Partner channel: 2 ERP partners",
               "Retainer #2; automation of weekly pack", "Review: kill / double-down decision"]
    for w in range(13):
        r = w + 2
        t.cell(row=r, column=1, value=w + 1)
        t.cell(row=r, column=2, value=phases[w])
        t.cell(row=r, column=3, value=targets[w])
        for c in range(4, 12):
            t.cell(row=r, column=c).fill = INP
    t.cell(row=15, column=1, value="Total").font = Font(bold=True)
    for c in range(4, 12):
        L = get_column_letter(c)
        t.cell(row=15, column=c, value=f"=SUM({L}2:{L}14)")
    sv = DataValidation(type="list", formula1='"On track,At risk,Off track,Done"', allow_blank=True)
    t.add_data_validation(sv)
    sv.add("L2:L14")
    t.conditional_formatting.add("L2:L14", CellIsRule(operator="equal", formula=['"Off track"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    t.conditional_formatting.add("L2:L14", CellIsRule(operator="equal", formula=['"On track"'], fill=PatternFill("solid", fgColor="C6EFCE")))

    # KPI dashboard
    d = wb.create_sheet("KPI Dashboard", 0)
    header(d, 1, ["KPI", "Formula / source", "Target (Day 90)", "Actual", "RAG"], [34, 60, 16, 14, 10])
    kpis = [
        ("Prospects in CRM", "=COUNTA('CRM Pipeline'!B2:B1000)", 400),
        ("Contacted", "=COUNTA('CRM Pipeline'!M2:M1000)-COUNTIF('CRM Pipeline'!M2:M1000,\"1*\")", 300),
        ("Reply rate %", "=IFERROR((COUNTIF('CRM Pipeline'!M2:M1000,\"3*\")+COUNTIF('CRM Pipeline'!M2:M1000,\"4*\")+COUNTIF('CRM Pipeline'!M2:M1000,\"5*\")+COUNTIF('CRM Pipeline'!M2:M1000,\"6*\"))/(COUNTA('CRM Pipeline'!M2:M1000)-COUNTIF('CRM Pipeline'!M2:M1000,\"1*\")),0)", 0.08),
        ("Discovery calls booked", "=COUNTIF('CRM Pipeline'!M2:M1000,\"4*\")+COUNTIF('CRM Pipeline'!M2:M1000,\"5*\")+COUNTIF('CRM Pipeline'!M2:M1000,\"6*\")", 12),
        ("Deals won", "=COUNTIF('CRM Pipeline'!M2:M1000,\"6*\")", 4),
        ("Weighted pipeline USD", "=SUM('CRM Pipeline'!U2:U1000)", 6000),
        ("Revenue to date USD (tracker)", "='90-Day Tracker'!J15", 3000),
        ("Effective hourly rate USD", "=IFERROR('90-Day Tracker'!J15/'90-Day Tracker'!K15,0)", 25),
        ("Base-case cumulative profit M3 PKR", "=Base!Q4", 0),
    ]
    for i, (n, fx, tgt) in enumerate(kpis, 2):
        d.cell(row=i, column=1, value=n)
        d.cell(row=i, column=2, value=fx.lstrip("="))
        d.cell(row=i, column=3, value=tgt)
        d.cell(row=i, column=4, value=fx)
        d.cell(row=i, column=5, value=f'=IF(D{i}>=C{i},"G",IF(D{i}>=0.7*C{i},"A","R"))')
    for r in (4,):
        d.cell(row=r, column=3).number_format = d.cell(row=r, column=4).number_format = "0.0%"
    d.conditional_formatting.add("E2:E20", CellIsRule(operator="equal", formula=['"G"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    d.conditional_formatting.add("E2:E20", CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", fgColor="FFEB9C")))
    d.conditional_formatting.add("E2:E20", CellIsRule(operator="equal", formula=['"R"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    d["A12"] = "Yellow cells on every sheet are inputs. Stage codes start with a digit (1-8); KPIs count stages with wildcard COUNTIF, e.g. \"3*\"."
    d["A12"].font = Font(italic=True, color="7F7F7F")

    wb.save(OUT / "Business_OS.xlsx")


if __name__ == "__main__":
    (OUT / "scoring_table.md").write_text(scoring_md() + "\n")
    (OUT / "financial_model.md").write_text(financial_md() + "\n")
    build_xlsx()
    print(scoring_md())
    print(financial_md())
